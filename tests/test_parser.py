import pytest
import os
import tempfile
from datetime import date, datetime
from openpyxl import Workbook

from app.parser import parse_orario, parse_assenze, parse_all, parse_visite
from app.models import Docente, Assenza


# ----------------------------------------------------------------------
# Helper per creare file Excel temporanei
# ----------------------------------------------------------------------

def crea_excel_orario(docenti_data, colonna_tipo=False):
    """
    Crea un file Excel di orario.
    docenti_data: lista di elementi che possono essere:
      - tuple (nome, lista_valori)  -> riga con nome + valori
      - tuple (None, lista_valori)  -> riga vuota (solo per test)
      - lista già piatta            -> riga così com'è
    """
    wb = Workbook()
    ws = wb.active

    headers = ["Docente"]
    if colonna_tipo:
        headers.append("Tipo")
    headers += [
        f"{giorno}-{ora}"
        for giorno in ["Lun", "Mar", "Mer", "Gio", "Ven"]
        for ora in range(1, 8)
    ]
    ws.append(headers)

    for item in docenti_data:
        if isinstance(item, tuple) and len(item) == 2:
            nome, valori = item
            ws.append([nome] + list(valori))
        else:
            ws.append(item)

    temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
    wb.save(temp.name)
    temp.close()
    return temp.name


def crea_excel_assenze(assenze_data):
    """Crea un file Excel di assenze.
    assenze_data: lista di tuple (docente, data, ora_inizio, ora_fine, tipo, data_com)
    """
    wb = Workbook()
    ws = wb.active

    headers = ["Docente", "Data", "Ora_inizio", "Ora_fine", "Tipo_assenza", "Data_comunicazione"]
    ws.append(headers)

    for ass in assenze_data:
        riga = list(ass)
        for i in [1, 5]:
            if isinstance(riga[i], date):
                riga[i] = datetime.combine(riga[i], datetime.min.time())
        ws.append(riga)

    temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
    wb.save(temp.name)
    temp.close()
    return temp.name


def crea_excel_visite(visite_data):
    """Crea un file Excel di visite (solo per test)."""
    wb = Workbook()
    ws = wb.active
    headers = ["Classe", "Data", "Docenti_accompagnatori"]
    ws.append(headers)
    for vis in visite_data:
        ws.append([vis[0], vis[1], vis[2]])
    temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
    wb.save(temp.name)
    temp.close()
    return temp.name


# ----------------------------------------------------------------------
# Test parse_orario
# ----------------------------------------------------------------------

class TestParseOrario:
    def test_orario_valido(self):
        dati = [
            ("Rossi", ["1A", "1A", None, "D", None, None, None,
                       "2B", None, None, None, None, None, None,
                       None, None, None, None, None, None, None,
                       None, None, None, None, None, None, None,
                       "1A", "1A", "D", None, None, None, None]),
        ]
        file_path = crea_excel_orario(dati)
        try:
            docenti = parse_orario(file_path)
            assert len(docenti) == 1
            docente = docenti[0]
            assert docente.nome == "Rossi"
            assert docente.tipo == "materia"
            assert docente.orario[(1, 1)] == "1A"
            assert docente.orario[(1, 2)] == "1A"
            assert docente.orario[(1, 3)] is None
            assert docente.orario[(1, 4)] == "D"
            assert docente.orario[(2, 1)] == "2B"
            assert docente.orario[(5, 1)] == "1A"
            assert docente.orario[(5, 3)] == "D"
        finally:
            os.unlink(file_path)

    def test_orario_con_righe_vuote(self):
        dati = [
            ("Rossi", ["1A"] + [None] * 34),
            (None, [None] * 35),
        ]
        file_path = crea_excel_orario(dati)
        try:
            docenti = parse_orario(file_path)
            assert len(docenti) == 1
            assert docenti[0].nome == "Rossi"
        finally:
            os.unlink(file_path)

    def test_orario_missing_file(self):
        with pytest.raises(FileNotFoundError):
            parse_orario("file_inesistente.xlsx")

    def test_orario_intestazione_errata(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Nome", "Colonna1", "Colonna2"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            with pytest.raises(ValueError, match="La prima colonna deve essere 'Docente'"):
                parse_orario(temp.name)
        finally:
            os.unlink(temp.name)

    def test_orario_giorno_non_riconosciuto(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Docente", "Xxx-1", "Lun-2"])
        ws.append(["Rossi", "1A", "1A"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            docenti = parse_orario(temp.name)
            assert len(docenti) == 1
            assert (1, 2) in docenti[0].orario
            assert (1, 1) not in docenti[0].orario
        finally:
            os.unlink(temp.name)

    def test_orario_colonne_extra(self):
        wb = Workbook()
        ws = wb.active
        headers = ["Docente"] + [f"{g}-{o}" for g in ["Lun", "Mar", "Mer", "Gio", "Ven"] for o in range(1, 8)] + ["Extra1", "Extra2"]
        ws.append(headers)
        ws.append(["Rossi"] + ["1A"] * 35 + ["ignorato", "ignorato"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            docenti = parse_orario(temp.name)
            assert len(docenti) == 1
            assert len(docenti[0].orario) == 35
        finally:
            os.unlink(temp.name)

    def test_orario_valori_non_stringa(self):
        wb = Workbook()
        ws = wb.active
        headers = ["Docente"] + [f"{g}-{o}" for g in ["Lun", "Mar", "Mer", "Gio", "Ven"] for o in range(1, 8)]
        ws.append(headers)
        ws.append(["Rossi"] + [123] + [None] * 34)
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            docenti = parse_orario(temp.name)
            assert docenti[0].orario[(1, 1)] == "123"
        finally:
            os.unlink(temp.name)

    def test_orario_spazi_bianchi(self):
        wb = Workbook()
        ws = wb.active
        headers = ["Docente"] + [f"{g}-{o}" for g in ["Lun", "Mar", "Mer", "Gio", "Ven"] for o in range(1, 8)]
        ws.append(headers)
        ws.append(["Rossi"] + ["", "   "] + [None] * 33)
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            docenti = parse_orario(temp.name)
            assert docenti[0].orario[(1, 1)] is None
            assert docenti[0].orario[(1, 2)] is None
        finally:
            os.unlink(temp.name)

    def test_orario_intestazione_con_spazi(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Docente", "Lun-1 ", "Lun-2"])
        ws.append(["Rossi", "1A", "1A"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            docenti = parse_orario(temp.name)
            assert len(docenti) == 1
            assert docenti[0].orario[(1, 1)] == "1A"
            assert docenti[0].orario[(1, 2)] == "1A"
        finally:
            os.unlink(temp.name)

    # --- US18: colonna Tipo e compresenze ---

    def test_orario_con_tipo_e_compresenza(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Docente", "Tipo", "Lun-1", "Lun-2"])
        ws.append(["Rossi", "materia", "1A", "D"])
        ws.append(["ITP1", "itp", "1A:Rossi", "D"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            docenti = parse_orario(temp.name)
            assert len(docenti) == 2
            rossi = next(d for d in docenti if d.nome == "Rossi")
            itp = next(d for d in docenti if d.nome == "ITP1")
            assert rossi.tipo == "materia"
            assert itp.tipo == "itp"
            assert itp.is_compresenza(1, 1) is True
            assert itp.get_classe_compresenza(1, 1) == "1A"
            assert itp.get_docente_materia_compresenza(1, 1) == "Rossi"
        finally:
            os.unlink(temp.name)

    def test_compresenza_con_docente_materia_inesistente(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Docente", "Tipo", "Lun-1"])
        ws.append(["ITP1", "itp", "1A:NonEsiste"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            with pytest.raises(ValueError, match="non presente nell'orario"):
                parse_orario(temp.name)
        finally:
            os.unlink(temp.name)

    def test_compresenza_con_docente_non_materia(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Docente", "Tipo", "Lun-1"])
        ws.append(["ITP1", "itp", "1A:ITP2"])
        ws.append(["ITP2", "itp", "D"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            with pytest.raises(ValueError, match="non è un docente di materia"):
                parse_orario(temp.name)
        finally:
            os.unlink(temp.name)

    def test_tipo_non_valido(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Docente", "Tipo", "Lun-1"])
        ws.append(["Rossi", "strano", "1A"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            with pytest.raises(ValueError, match="Tipo non valido"):
                parse_orario(temp.name)
        finally:
            os.unlink(temp.name)

    # tests/test_parser.py
# ... contenuto esistente ...

class TestParseOrarioCompresenze:
    def test_compresenza_con_docente_materia_inesistente(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Docente", "Tipo", "Lun-1"])
        ws.append(["ITP1", "itp", "1A:NonEsiste"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            with pytest.raises(ValueError, match="non presente nell'orario"):
                parse_orario(temp.name)
        finally:
            os.unlink(temp.name)

    def test_compresenza_con_docente_non_materia(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Docente", "Tipo", "Lun-1"])
        ws.append(["ITP1", "itp", "1A:ITP2"])
        ws.append(["ITP2", "itp", "D"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            with pytest.raises(ValueError, match="non è un docente di materia"):
                parse_orario(temp.name)
        finally:
            os.unlink(temp.name)

    def test_itp_con_codice_classe_senza_compresenza(self):
        """Un ITP con cella '1A' (senza ':') deve essere rifiutato."""
        wb = Workbook()
        ws = wb.active
        ws.append(["Docente", "Tipo", "Lun-1"])
        ws.append(["Rossi", "materia", "1A"])
        ws.append(["ITP1", "itp", "1A"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            with pytest.raises(ValueError, match="senza formato compresenza"):
                parse_orario(temp.name)
        finally:
            os.unlink(temp.name)

    def test_inferenza_tipo_da_prefissi(self):
        """Senza colonna Tipo, ITP_ e Sost_ inferiscono il tipo."""
        wb = Workbook()
        ws = wb.active
        ws.append(["Docente", "Lun-1", "Lun-2"])
        ws.append(["Rossi", "1A", "D"])
        ws.append([])  # riga vuota separatore
        ws.append(["ITP_Mario", "1A:Rossi", "D"])
        ws.append(["Sost_Anna", "1A:Rossi", "D"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            docenti = parse_orario(temp.name)
            tipi = {d.nome: d.tipo for d in docenti}
            assert tipi["Rossi"] == "materia"
            assert tipi["Mario"] == "itp"
            assert tipi["Anna"] == "sostegno"
            assert len(docenti) == 3
        finally:
            os.unlink(temp.name)

# ----------------------------------------------------------------------
# Test parse_assenze
# ----------------------------------------------------------------------

class TestParseAssenze:
    def test_assenze_valide(self):
        dati = [
            ("Rossi", date(2025, 3, 14), 1, 2, "PERMESSO", date(2025, 3, 13)),
            ("Bianchi", date(2025, 3, 15), 3, 5, "MALATTIA", date(2025, 3, 14)),
        ]
        file_path = crea_excel_assenze(dati)
        try:
            assenze = parse_assenze(file_path)
            assert len(assenze) == 2
            a1 = assenze[0]
            assert a1.docente_nome == "Rossi"
            assert a1.data == date(2025, 3, 14)
            assert a1.ora_inizio == 1
            assert a1.ora_fine == 2
            assert a1.tipo_assenza == "PERMESSO"
            assert a1.data_comunicazione == date(2025, 3, 13)
        finally:
            os.unlink(file_path)

    def test_assenze_ora_inizio_fuori_range(self):
        dati = [("Rossi", date(2025, 3, 14), 0, 2, "PERMESSO", date(2025, 3, 13))]
        file_path = crea_excel_assenze(dati)
        try:
            with pytest.raises(ValueError, match="Ora_inizio deve essere compreso tra 1 e 7"):
                parse_assenze(file_path)
        finally:
            os.unlink(file_path)

    def test_assenze_ora_fine_minore_inizio(self):
        dati = [("Rossi", date(2025, 3, 14), 3, 2, "PERMESSO", date(2025, 3, 13))]
        file_path = crea_excel_assenze(dati)
        try:
            with pytest.raises(ValueError, match="Ora_fine non può essere minore di Ora_inizio"):
                parse_assenze(file_path)
        finally:
            os.unlink(file_path)

    def test_assenze_intestazione_errata(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Docente", "Data", "Ora_inizio", "Ora_fine"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            with pytest.raises(ValueError, match="Intestazioni non corrispondenti"):
                parse_assenze(temp.name)
        finally:
            os.unlink(temp.name)

    def test_assenze_ora_come_float(self):
        dati = [("Rossi", date(2025, 3, 14), 2.0, 3.0, "PERMESSO", date(2025, 3, 13))]
        file_path = crea_excel_assenze(dati)
        try:
            assenze = parse_assenze(file_path)
            assert len(assenze) == 1
            assert assenze[0].ora_inizio == 2
            assert assenze[0].ora_fine == 3
        finally:
            os.unlink(file_path)

    def test_assenze_ora_come_float_non_intero(self):
        dati = [("Rossi", date(2025, 3, 14), 2.5, 3, "PERMESSO", date(2025, 3, 13))]
        file_path = crea_excel_assenze(dati)
        try:
            with pytest.raises(ValueError, match="Ora_inizio deve essere un numero intero"):
                parse_assenze(file_path)
        finally:
            os.unlink(file_path)

    def test_assenze_data_come_stringa(self):
        wb = Workbook()
        ws = wb.active
        headers = ["Docente", "Data", "Ora_inizio", "Ora_fine", "Tipo_assenza", "Data_comunicazione"]
        ws.append(headers)
        ws.append(["Rossi", "2025-03-14", 2, 3, "PERMESSO", "2025-03-13"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            assenze = parse_assenze(temp.name)
            assert len(assenze) == 1
            a = assenze[0]
            assert a.docente_nome == "Rossi"
            assert a.data == date(2025, 3, 14)
            assert a.ora_inizio == 2
            assert a.ora_fine == 3
            assert a.tipo_assenza == "PERMESSO"
            assert a.data_comunicazione == date(2025, 3, 13)
        finally:
            os.unlink(temp.name)

    def test_assenze_riga_vuota(self):
        dati = [
            ("Rossi", date(2025, 3, 14), 1, 2, "PERMESSO", date(2025, 3, 13)),
            (None, None, None, None, None, None),
        ]
        file_path = crea_excel_assenze(dati)
        try:
            assenze = parse_assenze(file_path)
            assert len(assenze) == 1
        finally:
            os.unlink(file_path)

    def test_assenze_docente_non_stringa(self):
        dati = [(123, date(2025, 3, 14), 1, 2, "PERMESSO", date(2025, 3, 13))]
        file_path = crea_excel_assenze(dati)
        try:
            with pytest.raises(ValueError, match="Nome docente deve essere una stringa"):
                parse_assenze(file_path)
        finally:
            os.unlink(file_path)

    def test_assenze_tipo_assenza_numerico(self):
        dati = [("Rossi", date(2025, 3, 14), 1, 2, 123, date(2025, 3, 13))]
        file_path = crea_excel_assenze(dati)
        try:
            with pytest.raises(ValueError, match="Tipo_assenza deve essere una stringa"):
                parse_assenze(file_path)
        finally:
            os.unlink(file_path)

    def test_assenze_colonne_mancanti(self):
        wb = Workbook()
        ws = wb.active
        headers = ["Docente", "Data", "Ora_inizio", "Ora_fine", "Tipo_assenza"]
        ws.append(headers)
        ws.append(["Rossi", date(2025, 3, 14), 1, 2, "PERMESSO"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            with pytest.raises(ValueError, match="Intestazioni non corrispondenti"):
                parse_assenze(temp.name)
        finally:
            os.unlink(temp.name)

    def test_assenze_data_comunicazione_mancante(self):
        dati = [("Rossi", date(2025, 3, 14), 1, 2, "PERMESSO", None)]
        file_path = crea_excel_assenze(dati)
        try:
            with pytest.raises(ValueError, match="Data comunicazione non valida"):
                parse_assenze(file_path)
        finally:
            os.unlink(file_path)


# ----------------------------------------------------------------------
# Test parse_visite
# ----------------------------------------------------------------------

class TestParseVisite:
    def test_visite_valide(self):
        dati = [
            ("1A", date(2025, 3, 14), "Rossi, Verdi"),
            ("2B", date(2025, 3, 15), "Bianchi"),
        ]
        file_path = crea_excel_visite(dati)
        try:
            visite = parse_visite(file_path, ore_giornaliere=6)
            assert len(visite) == 2
            v1 = visite[0]
            assert v1.classe == "1A"
            assert v1.data == date(2025, 3, 14)
            assert v1.ora_inizio == 1
            assert v1.ora_fine == 6
            assert v1.docenti_accompagnatori == ["Rossi", "Verdi"]
            v2 = visite[1]
            assert v2.classe == "2B"
            assert v2.docenti_accompagnatori == ["Bianchi"]
        finally:
            os.unlink(file_path)

    def test_visite_intestazione_errata(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Classe", "Data", "Ora_inizio", "Docenti"])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            with pytest.raises(ValueError, match="Intestazioni non corrispondenti"):
                parse_visite(temp.name)
        finally:
            os.unlink(temp.name)

    def test_visite_docenti_mancanti(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Classe", "Data", "Docenti_accompagnatori"])
        ws.append(["1A", datetime(2025, 3, 14), None])
        temp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp.name)
        temp.close()
        try:
            visite = parse_visite(temp.name)
            assert len(visite) == 1
            assert visite[0].docenti_accompagnatori == []
        finally:
            os.unlink(temp.name)


# ----------------------------------------------------------------------
# Test parse_all
# ----------------------------------------------------------------------

class TestParseAll:
    def test_parse_all_success(self):
        orario_dati = [("Rossi", ["1A"] * 7 + [None] * 28)]
        orario_path = crea_excel_orario(orario_dati)

        assenze_dati = [("Rossi", date(2025, 3, 14), 1, 2, "PERMESSO", date(2025, 3, 13))]
        assenze_path = crea_excel_assenze(assenze_dati)

        try:
            docenti, assenze, visite = parse_all(orario_path, assenze_path, None, 6)
            assert len(docenti) == 1
            assert len(assenze) == 1
            assert len(visite) == 0
        finally:
            os.unlink(orario_path)
            os.unlink(assenze_path)